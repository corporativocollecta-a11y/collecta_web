"""
Vista Nueva Zelanda: producción por región (16 consejos regionales). Stats NZ publica por región la SUPERFICIE
de frutales y hortalizas, no la producción: la producción nacional de FAOSTAT 2024 se reparte según la superficie
oficial de cada región (todo marcado como estimado). Población estimada por región de Stats NZ.

Fuentes (Stats NZ, https://www.stats.govt.nz/information-releases/…):
  - Agricultural production statistics: Year to June 2024 (final), tablas 9 y 10: superficie plantada de frutales
    (manzana, aguacate, kiwi, uva de vino, cereza) y superficie cosechada de hortalizas (cebolla, papa, calabaza
    buttercup) por región. aps2024.xlsx. "S" = cifra suprimida por confidencialidad.
  - Agricultural production statistics: Year to June 2022 (final), resultados por región del censo agropecuario 2022
    (tabla 1, más de 40 cultivos hortícolas, con ruido aleatorio por confidencialidad). aps2022_regional.xlsx.
  - Subnational population estimates: At 30 June 2025 (provisional), tabla 1 (consejos regionales).
    poblacion_2025.xlsx.
  - FAOSTAT 2024: producción nacional; cebolla (partida 402), cereza, kiwi y col del archivo bruto de FAOSTAT.
Clave de reparto: superficie 2024 cuando la encuesta 2024 la publica (las regiones con "S" reciben el resto del total
nacional según su superficie del censo 2022); si no, superficie del censo 2022. Tomate: superficie al aire libre más
invernadero ponderada por un rendimiento de referencia (invernadero ≈ 6 veces más por hectárea).
Durazno y mandarina se omiten: la cifra nacional de FAOSTAT no es verosímil (96 t y 545 t sobre ≈600 ha).

Uso:
  python scripts/procesar_statsnz.py
Salida: data/sub_nz.js → window.SUBNACIONAL.NZ
"""
import unicodedata
import warnings
from pathlib import Path

import openpyxl

from subnacional_comun import escribir_pais, fao_bruto, produccion_fao, resumen

warnings.filterwarnings("ignore")
RAIZ = Path(__file__).resolve().parent.parent
F = RAIZ / "data" / "fuentes" / "nz"
M49 = "554"

REGIONES = {  # código de consejo regional de Stats NZ → (nombre, lat, lon de la sede)
    "01": ("Northland", -35.73, 174.32), "02": ("Auckland", -36.85, 174.76), "03": ("Waikato", -37.79, 175.28),
    "04": ("Bahía de Plenty", -37.95, 176.99), "05": ("Gisborne", -38.66, 178.02), "06": ("Hawke's Bay", -39.49, 176.91),
    "07": ("Taranaki", -39.06, 174.08), "08": ("Manawatū-Whanganui", -40.35, 175.61), "09": ("Wellington", -41.29, 174.78),
    "12": ("Costa Oeste", -42.45, 171.21), "13": ("Canterbury", -43.53, 172.64), "14": ("Otago", -45.87, 170.50),
    "15": ("Southland", -46.41, 168.35), "16": ("Tasman", -41.34, 173.18), "17": ("Nelson", -41.27, 173.28),
    "18": ("Marlborough", -41.51, 173.96),
}
INGLES = {"01": "northland", "02": "auckland", "03": "waikato", "04": "bay of plenty", "05": "gisborne",
          "06": "hawke's bay", "07": "taranaki", "08": "manawatu-whanganui", "09": "wellington", "12": "west coast",
          "13": "canterbury", "14": "otago", "15": "southland", "16": "tasman", "17": "nelson", "18": "marlborough"}


def norm(t):
    t = unicodedata.normalize("NFD", str(t)).encode("ascii", "ignore").decode().lower().strip()
    for s in (" region", " regions"):
        if t.endswith(s):
            t = t[: -len(s)]
    return t.strip()


POR_NOMBRE = {v: k for k, v in INGLES.items()}

# producto → (columnas de la tabla 2024 [tabla, cultivo] o None, líneas del censo 2022 [código, factor])
# factor: ponderación de la superficie (m² de invernadero → ha, rendimiento relativo)
PRODUCTOS = {
    "manzana": (("Table 9", "Apples"), [("lc8500", 1)]),
    "aguacate": (("Table 9", "Avocados"), [("lc8570", 1)]),
    "kiwi": (("Table 9", "Kiwifruit"), [("lc8555", 1)]),
    "uva": (("Table 9", "Wine grapes"), [("lc8600", 1), ("lc8605", 1)]),
    "cereza": (("Table 9", "Cherries"), [("lc8540", 1)]),
    "cebolla": (("Table 10", "Onions"), [("lc9120", 1)]),
    "papa": (("Table 10", "Potatoes"), [("lc9135", 1)]),
    "pera": (None, [("lc8510", 1), ("lc8515", 1)]),
    "arandano": (None, [("lc8625", 1), ("lc8645", 1), ("lc8630", 1)]),
    "fresa": (None, [("lc8655", 1)]),
    "naranja": (None, [("lc8665", 1)]),
    "toronja": (None, [("lc8670", 1)]),
    "limon": (None, [("lc8675", 1)]),
    "esparrago": (None, [("lc9005", 1)]),
    "brocoli": (None, [("lc9025", 1), ("lc9050", 1)]),
    "col": (None, [("lc9035", 1)]),
    "zanahoria": (None, [("lc9045", 1)]),
    "lechuga": (None, [("lc9090", 1), ("lc9754", 1e-4)]),
    "melon": (None, [("lc9115", 1)]),
    "sandia": (None, [("lc9115", 1)]),
    "calabacita": (None, [("lc9140", 1), ("lc9170", 1)]),
    "jitomate": (None, [("lc9185", 1), ("lc9757", 6e-4)]),
    "chile": (None, [("lc9752", 1e-4)]),
    "pepino": (None, [("lc9753", 1e-4)]),
}
FAO_EXTRA = {"kiwi": "592", "cereza": "531", "col": "358", "cebolla": "402"}
NOMBRES = {"brocoli": "Brócoli y coliflor", "arandano": "Berries (arándano, frambuesa y boysenberry)",
           "chile": "Pimiento (invernadero)", "col": "Col (repollo)", "calabacita": "Calabaza y zapallo"}
NOTAS = {
    "kiwi": "Nueva Zelanda es el mayor exportador de kiwi del mundo; Bahía de Plenty concentra tres cuartas partes de los huertos.",
    "uva": "Casi toda la uva es para vino (Marlborough). Producción nacional de FAOSTAT repartida por superficie plantada.",
    "cebolla": "FAOSTAT registra la cebolla de Nueva Zelanda en la partida de cebolla verde (402); comercio de la partida de cebolla seca.",
    "jitomate": "Tomate al aire libre (industria, Hawke's Bay) e invernadero; la superficie de invernadero pesa 6 veces más por su rendimiento.",
    "lechuga": "Lechuga al aire libre y lechuga/hojas de ensalada de invernadero.",
    "melon": "Stats NZ publica sandía y melón juntos: ambos se reparten con la misma superficie.",
    "sandia": "Stats NZ publica sandía y melón juntos: ambos se reparten con la misma superficie.",
    "chile": "Solo pimiento de invernadero (superficie en m²).",
    "pepino": "Solo pepino de invernadero (superficie en m²).",
}


def num(v):
    return float(v) if isinstance(v, (int, float)) else None


def censo_2022():
    """{línea: {región: valor}} de la tabla 1 del censo 2022 (valores suprimidos = 0)."""
    ws = openpyxl.load_workbook(F / "aps2022_regional.xlsx", read_only=True).worksheets[1]
    filas = list(ws.iter_rows(values_only=True))
    lineas = filas[3]
    out = {}
    for r in filas[6:]:
        reg = POR_NOMBRE.get(norm(r[0] or ""))
        if not reg:
            continue
        for j, lc in enumerate(lineas):
            if lc and str(lc).startswith("lc"):
                out.setdefault(lc, {})[reg] = num(r[j]) or 0.0
    return out


def encuesta_2024():
    """{(tabla, cultivo): ({región: ha | None si 'S'}, total nacional)} de las tablas 9 y 10 de 2024."""
    wb = openpyxl.load_workbook(F / "aps2024.xlsx", data_only=True)
    out = {}
    for tabla in ("Table 9", "Table 10"):
        filas = list(wb[tabla].iter_rows(values_only=True))
        cultivos = None
        for r in filas:
            if r[0] is None and r[2] and isinstance(r[2], str) and not r[2].startswith("("):
                cultivos = {j: str(r[j]).strip() for j in range(2, len(r), 3) if r[j]}
                continue
            if cultivos is None or r[0] is None:
                continue
            nombre = norm(r[0])
            for j, c in cultivos.items():
                v = r[j + 1]   # columna 2024 (la anterior es 2022)
                d = out.setdefault((tabla, c), [{}, None])
                if nombre == "total new zealand":
                    d[1] = num(v)
                elif nombre in POR_NOMBRE:
                    d[0][POR_NOMBRE[nombre]] = None if v == "S" else (num(v) or 0.0)
    return out


def poblacion():
    ws = openpyxl.load_workbook(F / "poblacion_2025.xlsx", read_only=True)["Table 1"]
    pob = {}
    for r in ws.iter_rows(values_only=True):
        if r[0] and norm(r[0]) in POR_NOMBRE and isinstance(r[3], (int, float)):
            pob[POR_NOMBRE[norm(r[0])]] = int(r[3])   # 30 de junio de 2025 (provisional)
    return pob


def main():
    c22, e24, pob = censo_2022(), encuesta_2024(), poblacion()
    extra = fao_bruto(M49, FAO_EXTRA)
    extra_com = fao_bruto(M49, {"cebolla": "403"})["cebolla"]["comercio"]   # comercio de cebolla seca
    datos = {
        "codigo": "NZ", "m49": M49, "pais": "Nueva Zelanda", "nivel": "región", "nivelPlural": "regiones", "femenino": True,
        "anio": 2024, "anioPoblacion": 2025,
        "fuente": "Stats NZ Agricultural Production Survey 2024 y censo agropecuario 2022 (superficie por región), "
                  "Stats NZ población estimada 2025, FAOSTAT",
        "fuenteCorta": "Stats NZ · FAOSTAT",
        "metodo": "Stats NZ publica por región la superficie de frutales y hortalizas, no la producción: la producción "
                  "nacional de FAOSTAT 2024 se reparte según la superficie oficial de cada región (encuesta 2024 cuando existe; "
                  "censo agropecuario 2022 para los demás cultivos). Todo es estimado. Población estimada al 30 de junio de 2025.",
        "limites": [[-47.4, 166.3], [-34.3, 178.7]],
        "regiones": {c: {"nombre": n, "lat": la, "lon": lo, "pob": pob.get(c)} for c, (n, la, lo) in REGIONES.items()},
        "productos": {},
    }
    for clave, (col24, lineas) in PRODUCTOS.items():
        base22 = {c: sum(c22.get(lc, {}).get(c, 0) * f for lc, f in lineas) for c in REGIONES}
        if col24:
            reg24, total24 = e24[col24]
            areas = {c: v for c, v in reg24.items() if v is not None}
            supr = [c for c, v in reg24.items() if v is None]
            resto = max((total24 or 0) - sum(areas.values()), 0)
            peso = sum(base22[c] for c in supr)
            for c in supr:
                areas[c] = resto * (base22[c] / peso if peso else 1 / len(supr))
            for lc, f in lineas[1:]:   # uva de mesa: solo en el censo 2022
                for c in REGIONES:
                    areas[c] = areas.get(c, 0) + c22.get(lc, {}).get(c, 0) * f
            ha = total24
        else:
            areas = base22
            ha = sum(c22.get(lineas[0][0], {}).values())
        if clave in FAO_EXTRA:
            nacional = extra[clave]["prod"]
            com = extra_com if clave == "cebolla" else extra[clave]["comercio"]
        else:
            nacional, _ = produccion_fao(M49, clave)
            com = None
        total = sum(areas.values())
        if not nacional or not total:
            print(f"  (sin datos: {clave})")
            continue
        datos["productos"][clave] = {
            **({"nombre": NOMBRES[clave]} if clave in NOMBRES else {}),
            **({"nota": NOTAS[clave]} if clave in NOTAS else {}),
            "fuenteProduccion": "FAOSTAT (por superficie)", "anioProduccion": 2024,
            "nacional": round(nacional), "ha": round(ha or 0),
            "regiones": {c: [round(nacional * a / total), 1] for c, a in areas.items() if round(nacional * a / total) > 0},
            **({"comercio": com} if com else {}),
        }
    destino = escribir_pais(datos)
    print(f"-> {destino.name} ({destino.stat().st_size / 1e3:.0f} kB) · población {sum(pob.values()):,}")
    resumen(datos)


if __name__ == "__main__":
    main()
